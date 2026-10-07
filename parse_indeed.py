"""Indeed CSV -> indeed_求人一覧.xlsx. Appends only vacancies not already in the xlsx.

Usage: python3 parse_indeed.py [file1.csv file2.csv ...]
"""
import csv, re, sys, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

csv.field_size_limit(10**9)
OUT = 'indeed_求人一覧.xlsx'
srcs = sys.argv[1:] or ['dataset_indeed-scraper_2026-10-01_08-50-13-485.csv']
rows = [x for f in srcs for x in csv.DictReader(open(f, encoding='utf-8-sig'))]

PH = r'0\d{1,4}[-－‐−\s]?\d{1,4}[-－‐−\s]?\d{3,4}'
KW = r'(?:採用担当|問合せ先|問い合わせ先|お問い合わせ先|求人窓口|窓口|お問い合わせ電話番号|お問合せ電話番号|代表電話番号|電話番号|TEL|Tel|tel|ＴＥＬ|お電話|連絡先|電話)'


def key(u):
    m = re.search(r'jk=([0-9a-f]+)', u or '')
    return m.group(1) if m else (u or '').strip()


def norm(s):
    return s.translate(str.maketrans('０１２３４５６７８９－ー‐−', '0123456789----'))


def phone(d):
    d = norm(d)
    for m in re.finditer(KW + r'[^\d]{0,25}?[:：/]?\s*[【\(（]?(' + PH + r')', d):
        p = re.sub(r'\s', '', m.group(1))
        if re.fullmatch(r'0\d{1,4}-\d{1,4}-\d{3,4}', p):
            return p
    for m in re.finditer(r'【(' + PH + r')】', d):
        return m.group(1)
    return ''


def clean(s):
    return re.sub(r'^[\s:：・●■\*]+|[\s\*]+$', '', s)


def title(d):
    m = re.search(r'(?:募集職種|職種)\s*[:：]\s*([^\n]+)', d)
    if m and clean(m.group(1)):
        return clean(m.group(1))
    m = re.search(r'(?:募集職種|職種)\s*\n+\s*([^\n]+)', d)
    if m:
        t = clean(m.group(1))
        if t and not re.match(r'(職種解説|仕事内容|住所|所在地|給与)', t) and len(t) < 60:
            return t
    return ''


def addr(d, loc):
    m = re.search(r'(?:住所|所在地|勤務地)\s*[:：]?\s*\n*\s*(〒?\s*\d{3}-?\d{4}[^\n]*|(?:東京都|北海道|(?:京都|大阪)府|.{2,3}県)[^\n]{3,60})', d)
    return clean(m.group(1)) if m else loc


H = ['会社名', '役職', 'リンク(URL)', '時給', '電話番号', '住所']
if os.path.exists(OUT):
    wb = load_workbook(OUT)
    ws = wb.active
else:
    wb = Workbook()
    ws = wb.active
    ws.title = '求人'
    ws.append(H)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, w in enumerate([34, 40, 60, 26, 16, 45], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'

seen = {key(r[2].value) for r in ws.iter_rows(min_row=2) if r[2].value}
skip = dup = added = 0
for x in rows:
    d = x['description']
    if '派遣' in d + ''.join(v or '' for k, v in x.items() if k.startswith('jobType')) + (x.get('company') or ''):
        skip += 1
        continue
    k = key(x['url'])
    if k in seen:
        dup += 1
        continue
    seen.add(k)
    ws.append([x['company'], title(d), x['url'], x.get('salary', ''), phone(d), addr(d, x.get('location', ''))])
    c = ws.cell(ws.max_row, 3)
    c.hyperlink = c.value
    c.font = Font(color='0563C1', underline='single')
    added += 1

wb.save(OUT)
print(f'добавлено {added}, дубликатов пропущено {dup}, 派遣 пропущено {skip}, всего в файле {ws.max_row - 1}')
